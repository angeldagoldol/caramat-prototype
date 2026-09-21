#!/usr/bin/env python3
"""Build the Chapter I-II paper for the Online Scholarship Application System.

The document follows the SJPIIC/CICT Mini-System Project format: a cover page carrying the SJPIICD
logo, a table of contents, and Chapters I and II only. Chapter III onward is out of scope for this
submission and is deliberately not generated.

The system flowchart is drawn from code into ``assets/`` so the figure in the paper always matches
the prototype it documents.

Usage:
    python3 scripts/build_chapters_1_2_docx.py \
        --template assets/sjpiic-mini-system-format.docx \
        --logo assets/sjpiicd-logo.png \
        --output deliverables/Implementation_of_a_Priority_Queue_Max_Heap_in_an_Online_Scholarship_Application_System_Chapters_I_II.docx
"""
from __future__ import annotations

import argparse
import hashlib
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

TITLE = (
    "Implementation of a Priority Queue (Binary Max-Heap) in an "
    "Online Scholarship Application System"
)
FLOWCHART_ALT_TEXT = (
    "Flow from scholarship application entry through validation, merit scoring, max-heap insertion "
    "with sift-up, optional slot awarding through extract-max with sift-down, and display of the "
    "heap, the ranking, and the awarded slots."
)
FLOWCHART_PATH = ROOT / "assets" / "scholarship-system-flowchart.png"
DEFAULT_TEMPLATE = ROOT / "assets" / "sjpiic-mini-system-format.docx"
DEFAULT_LOGO = ROOT / "assets" / "sjpiicd-logo.png"
DEFAULT_OUTPUT = (
    ROOT
    / "deliverables"
    / "Implementation_of_a_Priority_Queue_Max_Heap_in_an_Online_Scholarship_Application_System_Chapters_I_II.docx"
)

NAVY = "0B2E59"
GOLD = "D4AF37"
SKY = "DCE8F5"
RED = "A63D40"
RED_SOFT = "FCE8E6"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the Chapter I-II scholarship system paper.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--logo", type=Path, default=DEFAULT_LOGO)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--flowchart-output", type=Path, default=FLOWCHART_PATH)
    return parser.parse_args()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reject_template_output_collision(template: Path, output: Path) -> None:
    """Refuse to write over the institutional template, however the paths are spelled."""
    template_resolved = template.resolve(strict=True)
    output_resolved = output.resolve(strict=False)
    same_resolved_path = template_resolved == output_resolved
    same_existing_file = output.exists() and os.path.samefile(template, output)
    if same_resolved_path or same_existing_file:
        raise ValueError("Refusing to overwrite the source template.")


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Arial"
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _apply_arial(style, size: float, *, bold: bool | None = None) -> None:
    style.font.name = "Arial"
    rpr = style._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold


def body_style_name(doc: Document) -> str:
    """The institutional template names its body style 'normal'; a blank document uses 'Normal'."""
    return "normal" if "normal" in doc.styles else "Normal"


def configure_styles(doc: Document) -> str:
    styles = doc.styles
    base_name = body_style_name(doc)

    normal = styles[base_name]
    _apply_arial(normal, 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    _apply_arial(title, 17, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(10)

    heading_1 = styles["Heading 1"]
    _apply_arial(heading_1, 17, bold=True)
    heading_1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_1.paragraph_format.first_line_indent = Inches(0)
    heading_1.paragraph_format.space_before = Pt(0)
    heading_1.paragraph_format.space_after = Pt(14)
    heading_1.paragraph_format.keep_with_next = True

    heading_2 = styles["Heading 2"]
    _apply_arial(heading_2, 13, bold=True)
    heading_2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading_2.paragraph_format.first_line_indent = Inches(0)
    heading_2.paragraph_format.space_before = Pt(11)
    heading_2.paragraph_format.space_after = Pt(7)
    heading_2.paragraph_format.keep_with_next = True

    def ensure(name: str):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]
        style.base_style = normal
        return style

    list_number = ensure("List Number")
    _apply_arial(list_number, 11)
    list_number.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    list_number.paragraph_format.left_indent = Inches(0.4)
    list_number.paragraph_format.first_line_indent = Inches(-0.28)
    list_number.paragraph_format.line_spacing = 1.15
    list_number.paragraph_format.space_after = Pt(5)

    cover = ensure("Cover Text")
    _apply_arial(cover, 11)
    cover.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.first_line_indent = Inches(0)
    cover.paragraph_format.line_spacing = 1.0
    cover.paragraph_format.space_after = Pt(3)

    toc_heading = ensure("TOC Heading")
    _apply_arial(toc_heading, 17, bold=True)
    toc_heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_heading.paragraph_format.first_line_indent = Inches(0)
    toc_heading.paragraph_format.space_after = Pt(16)

    toc_entry = ensure("TOC Entry")
    _apply_arial(toc_entry, 11)
    toc_entry.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc_entry.paragraph_format.first_line_indent = Inches(0)
    toc_entry.paragraph_format.line_spacing = 1.15
    toc_entry.paragraph_format.space_after = Pt(7)
    toc_entry.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    objective = ensure("Specific Objective")
    _apply_arial(objective, 11)
    objective.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    objective.paragraph_format.left_indent = Inches(0.4)
    objective.paragraph_format.first_line_indent = Inches(-0.28)
    objective.paragraph_format.line_spacing = 1.15
    objective.paragraph_format.space_after = Pt(4)

    caption = ensure("Figure Caption")
    _apply_arial(caption, 10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)

    return base_name


# ---------------------------------------------------------------------------
# Document plumbing
# ---------------------------------------------------------------------------


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    for pg_num in section._sectPr.findall(qn("w:pgNumType")):
        section._sectPr.remove(pg_num)


def clear_story(story) -> None:
    for paragraph in story.paragraphs:
        for child in list(paragraph._p):
            paragraph._p.remove(child)


def add_page_field(paragraph) -> None:
    """Insert a live PAGE field so Word renders the page number at the bottom of every page."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def restart_page_numbering(section) -> None:
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "decimal")
    pg_num.set(qn("w:start"), "1")
    section._sectPr.append(pg_num)


def set_alt_text(inline_shape, description: str, name: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("name", name)


def add_centered_picture(doc: Document, path: Path, width: float, alt_text: str, name: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_after = Pt(3)
    inline_shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(inline_shape, alt_text, name)
    return paragraph


def add_body_paragraph(doc: Document, text: str, *, lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(5)
    if lead:
        set_run_font(paragraph.add_run(lead), size=11, bold=True)
    set_run_font(paragraph.add_run(text), size=11)
    return paragraph


def add_numbered(doc: Document, number: int, text: str, style: str = "Specific Objective") -> None:
    paragraph = doc.add_paragraph(style=style)
    set_run_font(paragraph.add_run(f"{number}.  {text}"), size=11)
    paragraph.paragraph_format.keep_together = True


def add_cover(doc: Document, logo: Path) -> None:
    picture = add_centered_picture(
        doc, logo, 1.35, "Seal of St. John Paul II College of Davao.", "SJPIICD Logo"
    )
    picture.paragraph_format.space_after = Pt(7)

    institution = doc.add_paragraph("ST. JOHN PAUL II COLLEGE OF DAVAO", style="Cover Text")
    set_run_font(institution.runs[0], size=13, bold=True)
    cict = doc.add_paragraph("College of Information and Communications Technology", style="Cover Text")
    set_run_font(cict.runs[0], size=10)
    course = doc.add_paragraph("CC104 - Data Structures and Algorithms", style="Cover Text")
    set_run_font(course.runs[0], size=11)

    project = doc.add_paragraph("MINI-SYSTEM PROJECT", style="Cover Text")
    project.paragraph_format.space_before = Pt(10)
    project.paragraph_format.space_after = Pt(5)
    set_run_font(project.runs[0], size=12, bold=True)

    title = doc.add_paragraph(TITLE, style="Title")
    title.paragraph_format.keep_with_next = True

    for label, value in (
        ("Researchers", "[Full Names of Group Members]"),
        ("Group Leader", "[Full Name]"),
        ("Course", "Data Structures and Algorithms"),
        ("Instructor", "John Patrick P. Eleria"),
        ("Section", "[Section Name/Code]"),
        ("Institution", "St. John Paul II College of Davao"),
        ("Date of Submission", "[Month, Day, Year]"),
    ):
        label_paragraph = doc.add_paragraph(label, style="Cover Text")
        label_paragraph.paragraph_format.space_before = Pt(3)
        set_run_font(label_paragraph.runs[0], size=10, bold=True)
        value_paragraph = doc.add_paragraph(value, style="Cover Text")
        set_run_font(value_paragraph.runs[0], size=11)


def add_toc_entry(doc: Document, label: str, page: int, *, subsection: bool = False) -> None:
    paragraph = doc.add_paragraph(style="TOC Entry")
    if subsection:
        paragraph.paragraph_format.left_indent = Inches(0.3)
    set_run_font(paragraph.add_run(f"{label}\t{page}"), size=11)


def add_static_toc(doc: Document, pages: dict[str, int]) -> None:
    heading = doc.add_paragraph("TABLE OF CONTENTS", style="TOC Heading")
    heading.paragraph_format.space_before = Pt(3)
    for label, key, subsection in (
        ("Chapter I - Introduction", "chapter1", False),
        ("1.1 Background of the Study", "background", True),
        ("1.2 Objectives of the Study", "objectives", True),
        ("1.3 Significance of the Study", "significance", True),
        ("Chapter II - System Overview and Design", "chapter2", False),
        ("2.1 System Description", "description", True),
        ("2.2 System Architecture and Flowchart", "architecture", True),
        ("2.3 System Features", "features", True),
    ):
        add_toc_entry(doc, label, pages[key], subsection=subsection)

    note = doc.add_paragraph(style="TOC Entry")
    note.paragraph_format.space_before = Pt(12)
    set_run_font(
        note.add_run(
            "Chapters III to VI, References, and Appendices are not included in this submission. "
            "They are scheduled for the succeeding laboratory examinations."
        ),
        size=10,
        italic=True,
    )


# ---------------------------------------------------------------------------
# Chapter I
# ---------------------------------------------------------------------------


def add_chapter_one(doc: Document) -> None:
    heading = doc.add_paragraph("CHAPTER I - INTRODUCTION", style="Heading 1")
    heading.paragraph_format.page_break_before = True

    doc.add_paragraph("1.1 Background of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        "Scholarship programs in higher education institutions routinely receive far more "
        "applications than there are slots to award. A scholarship office therefore does not need "
        "the applications kept in full sorted order at all times. What it needs, repeatedly and on "
        "demand, is the answer to a single question: among the applications still waiting, who is "
        "the most deserving right now? Applications keep arriving while that question is being "
        "asked, so any approach that re-sorts the entire list after every new submission spends "
        "most of its work ordering applicants who will never be reached.",
    )
    add_body_paragraph(
        doc,
        "This study addresses that mismatch by building an Online Scholarship Application System "
        "around a priority queue implemented as a binary max-heap. A binary max-heap is a complete "
        "binary tree stored in a plain array, where the element at index i keeps its parent at index "
        "(i - 1) / 2 and its children at indices 2i + 1 and 2i + 2. The structure maintains one "
        "invariant, the heap property: every parent holds a priority greater than or equal to both "
        "of its children. The tree is therefore only partially ordered, which is precisely why it is "
        "inexpensive to maintain, yet the single element the scholarship office actually needs, the "
        "highest priority applicant, is always at the root and can be read in constant time.",
    )
    add_body_paragraph(
        doc,
        "In this system the priority of an application is a merit score from 0 to 100 that the "
        "system computes from three weighted components: academic standing from the general weighted "
        "average, worth up to 60 points; financial need from the reported monthly household income, "
        "worth up to 30 points; and the enrolled unit load for the term, worth up to 10 points. "
        "Because two applicants can earn exactly the same score, the queue breaks ties in favour of "
        "whoever submitted first, which keeps the ordering total, deterministic, and fair. Submitting "
        "an application is an insertion that sifts the new applicant up toward the root, and awarding "
        "a scholarship slot is an extract-max that removes the root and sifts its replacement back "
        "down. Both operations touch at most one node per level of the tree, so each costs O(log n) "
        "rather than the O(n log n) a full re-sort would cost.",
    )
    add_body_paragraph(
        doc,
        "The prototype implements the heap from first principles in a class named "
        "ApplicantPriorityQueue rather than delegating to the built-in java.util.PriorityQueue, so "
        "the index arithmetic, the sift-up loop, and the sift-down loop are all visible and can be "
        "explained line by line. The heap is the true core of the system: no sorting routine decides "
        "who is awarded a slot, and the ordering is never recomputed from scratch. To make the "
        "structure observable rather than merely claimed, the browser interface renders the live "
        "backing array in level order beside the tree it represents, and reports the number of "
        "comparisons and swaps each operation performed. Applications are held in the memory of the "
        "running Java process and are cleared when that process restarts.",
    )

    doc.add_paragraph("1.2 Objectives of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        "The general objective of this study is to implement a binary max-heap priority queue as the "
        "core mechanism of an Online Scholarship Application System, and to observe how that "
        "structure maintains the ordering of applicants as applications arrive and scholarship slots "
        "are awarded.",
    )
    add_body_paragraph(doc, "Specifically, this study aims:")
    for number, objective in enumerate(
        (
            "To implement the insert, extract-max, and peek operations of a binary max-heap from "
            "first principles over an array-backed complete binary tree, using only parent and child "
            "index arithmetic rather than stored node links.",
            "To convert each submitted application into a single merit score from 0 to 100 derived "
            "from academic standing, financial need, and enrolled unit load, and to settle equal "
            "scores by order of submission so the resulting ordering is total and deterministic.",
            "To enqueue every validated application through a sift-up so that the applicant holding "
            "the highest merit score always occupies the root of the heap.",
            "To award each scholarship slot through a single extract-max with a sift-down, so the "
            "most deserving waiting applicant is released without re-sorting the remaining "
            "applications.",
            "To expose the backing array, the tree it represents, the fully ranked waiting list, and "
            "the comparison and swap counts of each operation, so the behaviour of the heap can be "
            "inspected directly rather than inferred.",
            "To validate every application in Java before it reaches the structure, so that a "
            "rejected submission leaves the queue and the awarded slots unchanged.",
        ),
        start=1,
    ):
        add_numbered(doc, number, objective)

    doc.add_paragraph("1.3 Significance of the Study", style="Heading 2")
    add_body_paragraph(
        doc,
        "This study is significant to the following:",
    )
    add_body_paragraph(
        doc,
        " The project moves the priority queue from a definition to be memorised into a structure "
        "the researchers had to build, debug, and defend. Writing the sift-up and sift-down loops by "
        "hand and proving through testing that the heap property survives thousands of mixed "
        "operations taught the trade-off at the centre of the structure: a heap gives up full "
        "ordering in exchange for cheap access to the single element that matters.",
        lead="Researchers.",
    )
    add_body_paragraph(
        doc,
        " The interface shows the backing array and the tree side by side, which makes visible the "
        "point that is hardest to convey from a textbook diagram alone. A student can see that the "
        "array is genuinely not sorted, that only the root is guaranteed to be the maximum, and that "
        "an insertion still finds its place after only a handful of comparisons.",
        lead="Data Structures and Algorithms students.",
    )
    add_body_paragraph(
        doc,
        " The prototype offers a working example in which a priority queue is the core mechanism of "
        "a realistic system rather than a decorative add-on. The reported comparison and swap counts "
        "give an immediate, concrete demonstration of logarithmic growth that can be shown during a "
        "lecture without instrumentation.",
        lead="Instructors.",
    )
    add_body_paragraph(
        doc,
        " The system models the real constraint a scholarship committee works under: a fixed "
        "number of slots and a continuing stream of applicants. Because the merit score is broken "
        "down into its academic, financial-need, and unit-load components on screen, the ranking "
        "is explainable to an applicant who asks why another was selected first.",
        lead="Scholarship coordinators.",
    )
    add_body_paragraph(
        doc,
        " An applicant can see the exact basis of the ranking, including the position held in the "
        "queue and the points earned in each component. The submission-order tie-break also "
        "guarantees that applying earlier is never a disadvantage against an identically "
        "qualified candidate.",
        lead="Scholarship applicants.",
    )


# ---------------------------------------------------------------------------
# Chapter II
# ---------------------------------------------------------------------------


def add_chapter_two(doc: Document, flowchart: Path) -> None:
    heading = doc.add_paragraph("CHAPTER II - SYSTEM OVERVIEW AND DESIGN", style="Heading 1")
    heading.paragraph_format.page_break_before = True

    doc.add_paragraph("2.1 System Description", style="Heading 2")
    add_body_paragraph(
        doc,
        "The Online Scholarship Application System is a single Spring Boot application that both "
        "serves the browser interface and exposes the JSON endpoints behind it. An applicant is "
        "entered through a web form that collects five values: the applicant's name, the degree "
        "program, the general weighted average, the reported monthly household income, and the "
        "number of units enrolled for the term. The browser performs required-field checks first, "
        "but Java holds the authoritative validation, so a request that bypasses the browser is "
        "still rejected and the queue is left untouched.",
    )
    add_body_paragraph(
        doc,
        "Once an application is accepted, the system assigns it a submission sequence number and a "
        "reference code of the form SCH-0001, then computes its merit score. The score is the sum of "
        "three clamped components: a general weighted average of 1.00 earns the full 60 academic "
        "points and 3.00 or worse earns none; a monthly household income of zero earns the full 30 "
        "need points and an income of 60,000 pesos or more earns none; and 24 units or more earns "
        "the full 10 load points. That single number becomes the applicant's priority inside the "
        "queue, with the submission sequence number acting as the tie-breaker when two applicants "
        "score identically.",
    )
    add_body_paragraph(
        doc,
        "The scored applicant is then inserted into an ApplicantPriorityQueue, the hand-written "
        "binary max-heap that is the core of the system. The applicant is appended at the first free "
        "position of the backing array, which keeps the binary tree complete, and is then swapped "
        "upward with its parent for as long as it outranks that parent. When the scholarship office "
        "awards a slot, the system performs one extract-max: it takes the root, moves the last "
        "element into the root position, and sinks that element by repeatedly swapping it with its "
        "stronger child until the heap property holds again. The remaining applications are never "
        "re-sorted, and the number of scholarship slots is a configurable limit that defaults to "
        "five, after which further award requests are refused.",
    )
    add_body_paragraph(
        doc,
        "The interface presents the same data in three deliberately different views. The raw backing "
        "array is shown in level order with each element's index and parent index, making clear that "
        "the array is only partially ordered. A tree diagram renders the same array as the complete "
        "binary tree it represents, with the root highlighted. A separate ranked table lists every "
        "waiting applicant from highest to lowest merit, produced by running a heap sort over a copy "
        "of the array so that reading the ranking never disturbs the live queue. Each operation also "
        "reports the comparisons and swaps it performed. All applications are held in the memory of "
        "the running Java process, so the queue and the awarded slots reset whenever that process "
        "restarts.",
    )

    doc.add_paragraph("2.2 System Architecture and Flowchart", style="Heading 2")
    add_body_paragraph(
        doc,
        "The architecture keeps the browser responsible only for collecting input and presenting "
        "results, while Java owns validation, merit scoring, the heap itself, and the decision of who "
        "is awarded next. No ranking logic exists in JavaScript. Figure 1 shows the implemented flow, "
        "including the recoverable error path that retains the values the applicant typed and the "
        "branch taken when a scholarship slot is awarded.",
    )
    image_paragraph = add_centered_picture(
        doc, flowchart, 5.5, FLOWCHART_ALT_TEXT, "Online Scholarship Application System Flowchart"
    )
    image_paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph(
        "Figure 1. Online Scholarship Application System Flowchart", style="Figure Caption"
    )

    flow_input = add_body_paragraph(
        doc,
        " The applicant's name, degree program, general weighted average, monthly household income, "
        "and units enrolled, entered through the browser form. The general weighted average must "
        "fall from 1.00 through 5.00, the income from 0 through 1,000,000 pesos, and the unit load "
        "from 1 through 36. A second form of input is the award request raised by the scholarship "
        "office when a slot is to be granted.",
        lead="Input.",
    )
    flow_input.paragraph_format.page_break_before = True
    add_body_paragraph(
        doc,
        " The browser checks the required fields, then posts the application to Java. Java validates "
        "the fields and their ranges, assigns the submission sequence and reference code, and "
        "computes the merit score from its academic, financial-need, and unit-load components. The "
        "scored applicant is appended to the backing array and sifted up until its parent outranks "
        "it, restoring the heap property. When a slot is awarded instead, the system removes the "
        "root, moves the last element into its place, and sifts that element down until the heap "
        "property holds again, then records the awarded slot.",
        lead="Process.",
    )
    add_body_paragraph(
        doc,
        " The service returns the complete queue state as JSON: the backing array in level order, "
        "the fully ranked waiting list, the applicant currently at the root, the awarded slots, the "
        "remaining slot count, the height and capacity of the heap, and a re-check of the heap "
        "property. JavaScript renders the tree, the array table, the ranked table, the awarded table, "
        "and the comparison and swap counts of the operation just performed, then leaves the "
        "interface ready for the next action.",
        lead="Output.",
    )

    doc.add_paragraph("2.3 System Features", style="Heading 2")
    for number, feature in enumerate(
        (
            "Accepts a scholarship application consisting of an applicant name, degree program, "
            "general weighted average, monthly household income, and units enrolled, through a "
            "responsive browser form.",
            "Validates every field in Java, rejecting blank names, out-of-range averages, negative "
            "or excessive incomes, and invalid unit loads, while leaving the queue and the awarded "
            "slots unchanged and retaining the values the applicant typed.",
            "Computes a merit score from 0 to 100 for each application, combining up to 60 points "
            "for academic standing, up to 30 points for financial need, and up to 10 points for the "
            "enrolled unit load, and displays the three components beside the total.",
            "Assigns each accepted application a submission sequence number and a reference code of "
            "the form SCH-0001.",
            "Stores every waiting application in a hand-written binary max-heap held in an array, "
            "where a parent at index i keeps its children at indices 2i + 1 and 2i + 2.",
            "Inserts each accepted applicant with a sift-up so the highest merit applicant always "
            "occupies the root of the heap.",
            "Breaks an exact tie in merit score in favour of the applicant who submitted first, "
            "making the ordering total and deterministic.",
            "Awards a scholarship slot with a single extract-max and sift-down, releasing the "
            "highest merit waiting applicant without re-sorting the remaining applications.",
            "Enforces a configurable limit on scholarship slots, defaulting to five, and refuses "
            "further award requests once every slot has been granted or when no applications are "
            "waiting.",
            "Displays the live backing array in level order with each element's index and parent "
            "index, alongside a tree diagram of the same array with the root highlighted.",
            "Produces a fully ranked list of waiting applicants by running a heap sort over a copy "
            "of the array, so reading the ranking never modifies the live queue.",
            "Reports the number of comparisons and swaps performed by each insertion and extraction, "
            "together with a readable trace of every swap and the current height of the tree.",
            "Re-checks the heap property across the whole array after every operation and reports "
            "the result in the interface.",
            "Grows the backing array by doubling it when the tree fills it, so the queue is not "
            "limited by a capacity fixed in advance.",
            "Loads a fixed set of sample applications on request, including a late-arriving "
            "strongest applicant and an exact tie, so the structure can be demonstrated without "
            "typing, and clears the queue and awarded slots on request.",
        ),
        start=1,
    ):
        add_numbered(doc, number, feature, style="List Number")


# ---------------------------------------------------------------------------
# Flowchart
# ---------------------------------------------------------------------------


def load_chart_font(size: int, *, bold: bool = False):
    candidates = (
        "/usr/share/fonts/opentype/urw-base35/NimbusSans-Bold.otf"
        if bold
        else "/usr/share/fonts/opentype/urw-base35/NimbusSans-Regular.otf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_text(draw, box, text, font, *, fill, outline, radius=30, padding=70):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=6)
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - padding
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    block = "\n".join(lines)
    spacing = 10
    bbox = draw.multiline_textbbox((0, 0), block, font=font, spacing=spacing, align="center")
    x = (x1 + x2 - (bbox[2] - bbox[0])) / 2 - bbox[0]
    y = (y1 + y2 - (bbox[3] - bbox[1])) / 2 - bbox[1]
    draw.multiline_text((x, y), block, font=font, fill="#111827", spacing=spacing, align="center")


def draw_arrow(draw, start, end, *, color=f"#{NAVY}", width=9):
    draw.line((start, end), fill=color, width=width)
    x, y = end
    draw.polygon(((x, y), (x - 18, y - 28), (x + 18, y - 28)), fill=color)


def draw_polyline_arrow(draw, points, *, color=f"#{NAVY}", width=9):
    draw.line(points, fill=color, width=width, joint="curve")
    x, y = points[-1]
    px, py = points[-2]
    if y > py:
        head = ((x, y), (x - 18, y - 28), (x + 18, y - 28))
    elif y < py:
        head = ((x, y), (x - 18, y + 28), (x + 18, y + 28))
    elif x > px:
        head = ((x, y), (x - 28, y - 18), (x - 28, y + 18))
    else:
        head = ((x, y), (x + 28, y - 18), (x + 28, y + 18))
    draw.polygon(head, fill=color)


# Main-column nodes, drawn top to bottom. Steps that always run together are drawn as one node so
# the figure fits a single page at a readable size.
FLOW_NODES = (
    ("Start", "terminator"),
    ("Applicant enters name, program, GWA, income, and units", "input"),
    ("Browser checks the required fields, then POST /api/applications", "process"),
    ("Java validates each field and its range", "process"),
    ("Valid application?", "decision"),
    ("Compute merit score: academic 60 + need 30 + load 10", "process"),
    ("insert(): append at the last array index, then sift up while the applicant outranks its parent", "process"),
    ("Heap property restored; the root holds the highest merit", "process"),
    ("Award a scholarship slot?", "decision"),
    ("extractMax(): remove the root, move the last element up, then sift down while a child outranks it", "process"),
    ("Record the awarded slot and reduce the remaining slots", "process"),
    ("Return the queue state as JSON", "output"),
    ("Render the tree, array, ranking, awards, and operation counts", "output"),
    ("Await the next application or award", "terminator"),
)

VALIDATION_DECISION = 4
AWARD_DECISION = 8
REJOIN_NODE = 11  # "Return the queue state as JSON"


def build_flowchart(path: Path) -> None:
    width, height = 2500, 3390
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = load_chart_font(44)
    bold_font = load_chart_font(48, bold=True)
    small_font = load_chart_font(38, bold=True)

    main_left, main_right = 400, 1880
    main_centre = (main_left + main_right) // 2
    node_height, decision_height, gap = 160, 210, 70

    boxes = []
    y = 70
    for _, kind in FLOW_NODES:
        h = decision_height if kind == "decision" else node_height
        boxes.append((main_left, y, main_right, y + h))
        y += h + gap

    # Nodes
    for (label, kind), box in zip(FLOW_NODES, boxes):
        x1, y1, x2, y2 = box
        if kind == "decision":
            cx, cy = (x1 + x2) // 2, (y1 + y2) // 2
            points = ((cx, y1), (x2, cy), (cx, y2), (x1, cy))
            draw.polygon(points, fill=f"#{GOLD}")
            draw.line(points + (points[0],), fill=f"#{NAVY}", width=6, joint="curve")
            bbox = draw.textbbox((0, 0), label, font=bold_font)
            draw.text(
                (cx - (bbox[2] - bbox[0]) / 2, cy - (bbox[3] - bbox[1]) / 2 - 8),
                label,
                font=bold_font,
                fill="#111827",
            )
        else:
            rounded_text(
                draw,
                box,
                label,
                bold_font if kind == "terminator" else font,
                fill=f"#{GOLD}" if kind == "terminator" else f"#{SKY}",
                outline=f"#{NAVY}",
                radius=75 if kind == "terminator" else 28,
                padding=80,
            )

    # Straight arrows down the main column, skipping the two decision branches.
    for index in range(len(boxes) - 1):
        if index in (VALIDATION_DECISION, AWARD_DECISION):
            continue
        draw_arrow(draw, (main_centre, boxes[index][3]), (main_centre, boxes[index + 1][1]))

    # Validation decision: YES continues down, NO goes right to the reject box and loops back.
    decision = boxes[VALIDATION_DECISION]
    draw_arrow(draw, (main_centre, decision[3]), (main_centre, boxes[VALIDATION_DECISION + 1][1]))
    draw.text((main_centre + 25, decision[3] + 14), "YES", font=small_font, fill=f"#{NAVY}")

    decision_mid = (decision[1] + decision[3]) // 2
    reject_box = (2010, decision_mid - 165, 2450, decision_mid + 165)
    rounded_text(
        draw,
        reject_box,
        "Reject: show the errors and keep the entered values",
        font,
        fill=f"#{RED_SOFT}",
        outline=f"#{RED}",
        radius=28,
        padding=70,
    )
    draw_polyline_arrow(
        draw, [(decision[2], decision_mid), (reject_box[0], decision_mid)], color=f"#{RED}"
    )
    draw.text((decision[2] + 12, decision_mid - 88), "NO", font=small_font, fill=f"#{RED}")

    entry_box = boxes[1]
    entry_mid = (entry_box[1] + entry_box[3]) // 2
    draw_polyline_arrow(
        draw,
        [
            ((reject_box[0] + reject_box[2]) // 2, reject_box[1]),
            ((reject_box[0] + reject_box[2]) // 2, entry_mid),
            (main_right, entry_mid),
        ],
        color=f"#{RED}",
    )

    # Award decision: YES continues down, NO bypasses the extraction branch.
    award = boxes[AWARD_DECISION]
    draw_arrow(draw, (main_centre, award[3]), (main_centre, boxes[AWARD_DECISION + 1][1]))
    draw.text((main_centre + 25, award[3] + 14), "YES", font=small_font, fill=f"#{NAVY}")

    award_mid = (award[1] + award[3]) // 2
    rejoin = boxes[REJOIN_NODE]
    rejoin_mid = (rejoin[1] + rejoin[3]) // 2
    bypass_x = 2160
    draw_polyline_arrow(
        draw,
        [
            (award[2], award_mid),
            (bypass_x, award_mid),
            (bypass_x, rejoin_mid),
            (rejoin[2], rejoin_mid),
        ],
        color=f"#{NAVY}",
    )
    draw.text((award[2] + 12, award_mid - 88), "NO", font=small_font, fill=f"#{NAVY}")

    # Loop from the final terminator back to the entry node.
    await_box = boxes[-1]
    await_mid = (await_box[1] + await_box[3]) // 2
    draw_polyline_arrow(
        draw,
        [
            (await_box[0], await_mid),
            (210, await_mid),
            (210, entry_mid),
            (entry_box[0], entry_mid),
        ],
        color=f"#{GOLD}",
        width=11,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=True, dpi=(300, 300))


# ---------------------------------------------------------------------------
# Assembly
# ---------------------------------------------------------------------------


def build_document(template: Path, logo: Path, output: Path, flowchart_output: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    if not logo.is_file():
        raise FileNotFoundError(f"Logo not found: {logo}")
    reject_template_output_collision(template, output)
    reject_template_output_collision(template, flowchart_output)

    template_digest = sha256(template)
    build_flowchart(flowchart_output)

    doc = Document(template)
    clear_template_body(doc)
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.sections[0].footer.is_linked_to_previous = False
    clear_story(doc.sections[0].footer)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = (
        "CC104 Data Structures and Algorithms Mini-System Project, Chapters I and II"
    )
    doc.core_properties.keywords = (
        "priority queue, binary max-heap, scholarship application, Java, Spring Boot"
    )

    add_cover(doc, logo)

    content = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(content)
    content.header.is_linked_to_previous = False
    content.footer.is_linked_to_previous = False
    clear_story(content.header)
    clear_story(content.footer)
    restart_page_numbering(content)
    add_page_field(content.footer.paragraphs[0])

    add_static_toc(
        doc,
        {
            "chapter1": 2,
            "background": 2,
            "objectives": 2,
            "significance": 3,
            "chapter2": 4,
            "description": 4,
            "architecture": 4,
            "features": 6,
        },
    )

    add_chapter_one(doc)
    add_chapter_two(doc, flowchart_output)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    if sha256(template) != template_digest:
        raise RuntimeError("The source template changed during document generation.")


def main() -> None:
    args = parse_args()
    flowchart_output = args.flowchart_output.resolve()
    build_document(args.template.resolve(), args.logo.resolve(), args.output.resolve(), flowchart_output)
    print(f"Created {args.output}")
    print(f"Created {flowchart_output}")


if __name__ == "__main__":
    main()
