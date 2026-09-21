"""Structural checks on the generated Chapter I-II paper.

These assertions guard the parts of the SJPIIC/CICT format that are easy to break silently: the
cover page fields, the logo, the restart of page numbering, the PAGE field in the footer, the
presence of the flowchart with alt text, and the rule that this submission stops at Chapter II.

Run with:  python3 -m pytest tests/ -q
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import build_chapters_1_2_docx as builder  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> tuple[Document, Path]:
    out_dir = tmp_path_factory.mktemp("paper")
    output = out_dir / "paper.docx"
    flowchart = out_dir / "flowchart.png"
    builder.build_document(builder.DEFAULT_TEMPLATE, builder.DEFAULT_LOGO, output, flowchart)
    return Document(output), flowchart


@pytest.fixture(scope="module")
def doc(built) -> Document:
    return built[0]


@pytest.fixture(scope="module")
def texts(doc) -> list[str]:
    return [paragraph.text.strip() for paragraph in doc.paragraphs]


@pytest.fixture(scope="module")
def full_text(texts) -> str:
    return "\n".join(texts)


# --- cover page -------------------------------------------------------------


def test_cover_carries_every_required_field(texts):
    for required in (
        "ST. JOHN PAUL II COLLEGE OF DAVAO",
        "MINI-SYSTEM PROJECT",
        "Researchers",
        "Group Leader",
        "Course",
        "Instructor",
        "Section",
        "Institution",
        "Date of Submission",
    ):
        assert required in texts, f"cover page is missing {required!r}"


def test_cover_shows_the_full_title(texts):
    assert builder.TITLE in texts


def test_logo_is_the_first_image_and_has_alt_text(doc):
    shapes = doc.inline_shapes
    assert len(shapes) == 2, "expected exactly the logo and the flowchart"
    descr = shapes[0]._inline.docPr.get("descr")
    assert "St. John Paul II College of Davao" in descr


# --- table of contents ------------------------------------------------------


def test_table_of_contents_lists_only_chapters_one_and_two(full_text):
    assert "TABLE OF CONTENTS" in full_text
    for entry in (
        "Chapter I - Introduction",
        "1.1 Background of the Study",
        "1.2 Objectives of the Study",
        "1.3 Significance of the Study",
        "Chapter II - System Overview and Design",
        "2.1 System Description",
        "2.2 System Architecture and Flowchart",
        "2.3 System Features",
    ):
        assert entry in full_text, f"table of contents is missing {entry!r}"


def test_scope_note_states_the_later_chapters_are_excluded(full_text):
    assert "Chapters III to VI" in full_text
    assert "not included in this submission" in full_text


# --- chapters ---------------------------------------------------------------


def test_both_chapter_headings_are_present(texts):
    assert "CHAPTER I - INTRODUCTION" in texts
    assert "CHAPTER II - SYSTEM OVERVIEW AND DESIGN" in texts


def test_no_chapter_beyond_two_is_written(full_text):
    for forbidden in (
        "CHAPTER III",
        "CHAPTER IV",
        "CHAPTER V",
        "CHAPTER VI",
        "3.1 Overview of the Algorithm",
        "4.1 Time Complexity",
    ):
        assert forbidden not in full_text, f"{forbidden!r} is out of scope for this submission"


def test_every_required_subsection_heading_is_present(texts):
    for heading in (
        "1.1 Background of the Study",
        "1.2 Objectives of the Study",
        "1.3 Significance of the Study",
        "2.1 System Description",
        "2.2 System Architecture and Flowchart",
        "2.3 System Features",
    ):
        assert heading in texts


def test_chapters_start_on_a_new_page(doc):
    for heading in ("CHAPTER I - INTRODUCTION", "CHAPTER II - SYSTEM OVERVIEW AND DESIGN"):
        paragraph = next(p for p in doc.paragraphs if p.text.strip() == heading)
        assert paragraph.paragraph_format.page_break_before is True, f"{heading} must start a page"


def test_no_paragraph_exists_only_to_hold_a_page_break(doc):
    """A break-only paragraph can land on a fresh page and leave a blank page behind it."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            continue
        breaks = paragraph._p.findall(f".//{qn('w:br')}")
        page_breaks = [b for b in breaks if b.get(qn("w:type")) == "page"]
        assert not page_breaks, "use page_break_before instead of a break-only paragraph"


# --- content fidelity to the prototype --------------------------------------


def test_objectives_are_numbered_and_complete(full_text):
    for number in range(1, 7):
        assert f"{number}.  To " in full_text, f"objective {number} is missing"


def test_features_are_numbered_and_complete(full_text):
    for number in range(1, 16):
        assert f"{number}.  " in full_text, f"feature {number} is missing"


def test_significance_names_every_beneficiary(full_text):
    for beneficiary in (
        "Researchers.",
        "Data Structures and Algorithms students.",
        "Instructors.",
        "Scholarship coordinators.",
        "Scholarship applicants.",
    ):
        assert beneficiary in full_text


def test_input_process_output_are_each_discussed(full_text):
    for label in ("Input.", "Process.", "Output."):
        assert label in full_text


def test_document_describes_the_technique_the_prototype_actually_uses(full_text):
    for claim in (
        "binary max-heap",
        "ApplicantPriorityQueue",
        "sift-up",
        "sift-down",
        "O(log n)",
        "2i + 1",
        "2i + 2",
        "SCH-0001",
    ):
        assert claim in full_text, f"the paper should mention {claim!r}"


def test_scoring_weights_match_the_java_implementation(full_text):
    """These numbers must stay in step with MeritScorer."""
    assert "60 academic" in full_text or "up to 60 points" in full_text
    assert "30 need" in full_text or "up to 30 points" in full_text
    assert "10 load" in full_text or "up to 10 points" in full_text
    assert "60,000 pesos" in full_text
    assert "24 units" in full_text


# --- figure -----------------------------------------------------------------


def test_flowchart_is_embedded_with_alt_text_and_a_caption(doc, full_text):
    shape = doc.inline_shapes[1]
    assert shape._inline.docPr.get("descr") == builder.FLOWCHART_ALT_TEXT
    assert "Figure 1. Online Scholarship Application System Flowchart" in full_text


def test_flowchart_fits_one_portrait_page(built):
    """The figure plus its caption must not overflow the 9-inch content area."""
    from PIL import Image

    _, flowchart = built
    with Image.open(flowchart) as image:
        width_px, height_px = image.size
    rendered_height_in = 5.5 * height_px / width_px
    assert rendered_height_in < 8.5, f"figure renders {rendered_height_in:.2f}in tall"


# --- page setup -------------------------------------------------------------


def test_the_document_has_a_cover_section_and_a_numbered_body_section(doc):
    assert len(doc.sections) == 2

    cover, body = doc.sections
    assert cover._sectPr.findall(qn("w:pgNumType")) == []

    page_numbering = body._sectPr.findall(qn("w:pgNumType"))
    assert len(page_numbering) == 1
    assert page_numbering[0].get(qn("w:start")) == "1"


def test_page_numbers_are_a_live_field_in_the_footer(doc):
    footer = doc.sections[1].footer
    instructions = footer.paragraphs[0]._p.findall(f".//{qn('w:instrText')}")
    assert any("PAGE" in (node.text or "") for node in instructions)


def test_margins_and_page_size_follow_the_institutional_format(doc):
    from docx.shared import Inches

    for section in doc.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        for margin in (
            section.top_margin,
            section.bottom_margin,
            section.left_margin,
            section.right_margin,
        ):
            assert margin == Inches(1)


def test_body_text_is_arial_eleven_point(doc):
    normal = doc.styles[builder.body_style_name(doc)]
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 11


# --- safety -----------------------------------------------------------------


def test_the_builder_refuses_to_overwrite_the_template(tmp_path):
    with pytest.raises(ValueError, match="Refusing to overwrite"):
        builder.build_document(
            builder.DEFAULT_TEMPLATE,
            builder.DEFAULT_LOGO,
            builder.DEFAULT_TEMPLATE,
            tmp_path / "flowchart.png",
        )


def test_a_missing_logo_is_reported_clearly(tmp_path):
    with pytest.raises(FileNotFoundError, match="Logo not found"):
        builder.build_document(
            builder.DEFAULT_TEMPLATE,
            tmp_path / "absent.png",
            tmp_path / "paper.docx",
            tmp_path / "flowchart.png",
        )
